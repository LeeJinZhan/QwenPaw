"""S0 fixture for BR-038: synthetic multi-sheet workbook with merged cells.

Run with a Python environment that has openpyxl (e.g. agentic-runtime/.venv):
    ./.venv/bin/python tests/br038_s0/make_workbook.py <out_dir>

Produces workbook.xlsx plus ground_truth.json used by run_s0.py.
"""

from __future__ import annotations

import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook

SHEET_COUNT = 30
ROW_COUNT = 80
TEAM_SIZE = 10
COLUMNS = ["姓名", "岗位", "营销笔数", "营销金额(元)", "活跃天数", "备注"]


def build(out_dir: Path) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(20260919)
    workbook = Workbook()
    workbook.remove(workbook.active)
    truth: dict = {"sheets": {}}
    for index in range(1, SHEET_COUNT + 1):
        name = f"支行{index:02d}"
        sheet = workbook.create_sheet(name)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        sheet.cell(row=1, column=1, value=f"{name}2026年8月人员营销数据")
        for column, title in enumerate(COLUMNS, start=1):
            sheet.cell(row=2, column=column, value=title)
        sums = {"营销笔数": 0, "营销金额(元)": 0.0, "活跃天数": 0}
        for row in range(3, 3 + ROW_COUNT):
            team = (row - 3) // TEAM_SIZE
            if (row - 3) % TEAM_SIZE == 0:
                sheet.merge_cells(
                    start_row=row, start_column=2, end_row=row + TEAM_SIZE - 1, end_column=2
                )
                sheet.cell(row=row, column=2, value=f"团队{team + 1}")
            counts = rng.randint(1, 40)
            amount = round(rng.uniform(1000, 90000), 2)
            days = rng.randint(0, 31)
            sheet.cell(row=row, column=1, value=f"员工{row - 2:03d}")
            sheet.cell(row=row, column=3, value=counts)
            sheet.cell(row=row, column=4, value=amount)
            sheet.cell(row=row, column=5, value=days)
            sheet.cell(row=row, column=6, value="正常" if days else "待激活")
            sums["营销笔数"] += counts
            sums["营销金额(元)"] += amount
            sums["活跃天数"] += days
        truth["sheets"][name] = {
            "rows": ROW_COUNT,
            "sum_营销笔数": sums["营销笔数"],
            "sum_营销金额": round(sums["营销金额(元)"], 2),
            "sum_活跃天数": sums["活跃天数"],
        }
    target = out_dir / "workbook.xlsx"
    workbook.save(target)
    truth["sheet_count"] = SHEET_COUNT
    truth["row_count_per_sheet"] = ROW_COUNT
    truth["total_rows"] = SHEET_COUNT * ROW_COUNT
    (out_dir / "ground_truth.json").write_text(
        json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return truth


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/br038-s0")
    store_task = out / "store" / "taskS0"
    store_task.mkdir(parents=True, exist_ok=True)
    info = build(store_task)
    print(json.dumps({"workbook": str(store_task / "workbook.xlsx"), **{k: info[k] for k in ("sheet_count", "total_rows")}}, ensure_ascii=False))
