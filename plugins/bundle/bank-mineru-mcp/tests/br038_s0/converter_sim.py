"""S0 fixture for BR-038: simulate office-conversion markdown from the workbook.

Run with a Python environment that has openpyxl (e.g. agentic-runtime/.venv):
    ./.venv/bin/python tests/br038_s0/converter_sim.py <out_dir>

Emits md_full.md (all sheets) and md_trunc.md (first 8 sheets, standing in for a
source-side page/size cap). Merged cells are rendered the way a layout parser
sees them: value only in the anchor cell, blanks elsewhere, which reproduces the
"数值—列名对齐偏移" defect class without claiming MinerU-identical output.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

TRUNC_SHEETS = 8


def sheet_markdown(sheet) -> str:
    lines = [f"# {sheet.title}", ""]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return "\n".join(lines)
    width = max(len(row) for row in rows)
    for index, row in enumerate(rows):
        if index == 1:
            lines.append("| " + " | ".join(["---"] * width) + " |")
        cells = []
        for value in list(row) + [None] * (width - len(row)):
            cells.append("" if value is None else str(value))
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def main(out_dir: Path) -> None:
    workbook_path = out_dir / "store" / "taskS0" / "workbook.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheets = list(workbook.worksheets)
    full = "\n\n".join(sheet_markdown(sheet) for sheet in sheets)
    trunc = "\n\n".join(sheet_markdown(sheet) for sheet in sheets[:TRUNC_SHEETS])
    (out_dir / "md_full.md").write_text(full, encoding="utf-8")
    (out_dir / "md_trunc.md").write_text(trunc, encoding="utf-8")
    workbook.close()
    print(f"md_full={len(full)} chars, md_trunc={len(trunc)} chars, sheets={len(sheets)}")


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/br038-s0"))
