"""Capacity regressions exercise real files; large release fixtures run separately."""
import json
import tracemalloc
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from bank_mineru_mcp.spreadsheet import extract_workbook
from bank_mineru_mcp.document_store import DocumentStore
from bank_mineru_mcp.structured_store import StructuredStore
from bank_mineru_mcp.tools import MinerUToolService
from bank_runtime.gateway.document_access import approved_document_call, consume_document_call


def test_csv_extraction_has_bounded_memory(tmp_path):
    source = tmp_path / 'rows.csv'
    with source.open('w') as stream:
        stream.write('number,description\n')
        for i in range(18000):
            stream.write(f'{i},' + '建设计划' * 120 + '\n')
    tracemalloc.start()
    try:
        inventory = extract_workbook(source, tmp_path / 'out', stem='file')
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert inventory['sheets'][0]['rows'] == 18000
    assert peak < 12 * 1024 * 1024, f'whole-file buffering: {peak} bytes'


@pytest.mark.asyncio
async def test_many_merges_parse_keeps_usable_reference(tmp_path):
    task = tmp_path / 'task_test'
    task.mkdir()
    path = task / 'merges.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.append(['编号', '内容', '金额'])
    for i in range(2, 1202):
        ws.cell(i, 1, i)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        ws.cell(i, 3, 1)
    wb.save(path)
    source = SimpleNamespace(task_id='task_test', file_id='file_test', path=path,
        extension='.xlsx', media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    service = MinerUToolService(file_resolver=SimpleNamespace(resolve=lambda _: source),
        mineru_client=None, document_store=DocumentStore(root=tmp_path),
        structured_store=StructuredStore(root=tmp_path))
    with approved_document_call('task_test', 'parse_documents', {}) as meta:
        with consume_document_call(meta, 'parse_documents', {}):
            result = await service.parse_documents([{'file_id': 'file_test', 'file_ref': 'test-ref'}])
            assert result['status'] == 'completed'
            assert len(json.dumps(result, ensure_ascii=False, indent=2).encode()) <= 32000
            ref = result['items'][0]['document_ref']
            page = service.read_range(ref, rows=[1, 2])
            assert page['rows_scanned'] == 2


def test_numeric_merge_is_not_double_counted(tmp_path):
    task = tmp_path / 'task_test'
    task.mkdir()
    path = task / 'amounts.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.append(['项目', '金额'])
    ws.append(['A', 100])
    ws.append(['B', None])
    ws.merge_cells('B2:B3')
    wb.save(path)
    source = SimpleNamespace(task_id='task_test', path=path, expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    inventory = extract_workbook(path, task / 'work', stem='f')
    store = StructuredStore(root=tmp_path)
    handle = store.write(source, inventory, task/'work')
    result = store.aggregate(handle.document_ref, [{'metrics': [{'column':'金额', 'fn':'sum'}]}])
    assert result['results'][0]['groups'][0]['金额:sum'] == 100


@pytest.mark.asyncio
async def test_parse_job_reuses_published_result_after_restart(tmp_path):
    from bank_mineru_mcp.parse_jobs import parse_job
    task = tmp_path / 'task_test'
    task.mkdir()
    path = task / 'source.csv'
    path.write_text('id,value\n1,10\n2,20\n')
    source = SimpleNamespace(task_id='task_test', file_id='file_test', path=path,
        expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    store = StructuredStore(root=tmp_path)
    first, _ = await parse_job(store, source)
    restarted = StructuredStore(root=tmp_path)
    second, _ = await parse_job(restarted, source)
    assert first.document_ref == second.document_ref
    assert len(list((task/'.mineru-struct').glob('*/manifest.json'))) == 1
    path.write_text('id,value\n1,100\n')
    third, _ = await parse_job(restarted, source)
    assert third.document_ref != first.document_ref


def test_inventory_metadata_pages_preserve_all_merges(tmp_path):
    from bank_mineru_mcp.inventory import inventory_page
    meta = {'name':'data', 'columns':[{'name':'id'}], 'merged_ranges': [[i,1,i,2] for i in range(1200)]}
    inv = {'sheets':[meta]}
    cursor = 0
    entries = []
    while True:
        page = inventory_page(inv, sheet='data', start=cursor)
        assert len(json.dumps(page, ensure_ascii=False, indent=2).encode()) < 32000
        entries.extend(page['metadata'])
        if page['next_inventory_cursor'] is None:
            break
        assert page['next_inventory_cursor'] > cursor
        cursor = page['next_inventory_cursor']
    assert [x['range'] for x in entries if x['kind']=='merge'] == meta['merged_ranges']


def test_long_cell_has_lossless_fragment_access(tmp_path):
    task = tmp_path / 'task_test'
    task.mkdir()
    path = task / 'long.csv'
    original = '正文内容' * 10000
    path.write_text('id,text\n1,'+original+'\n')
    inventory = extract_workbook(path, task/'work', stem='f')
    source = SimpleNamespace(task_id='task_test',path=path,expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    store = StructuredStore(root=tmp_path)
    handle = store.write(source,inventory,task/'work')
    fragments,offset = [],0
    while True:
        page = store.read_cell(handle.document_ref,sheet='data',row=1,column='text',offset=offset)
        assert len(json.dumps(page,ensure_ascii=False,indent=2).encode()) < 32000
        fragments.append(page['text'])
        if page['next_cell_cursor'] is None:
            break
        offset=page['next_cell_cursor']
    assert ''.join(fragments)==original


@pytest.mark.asyncio
async def test_partial_formula_does_not_block_unrelated_statistics(tmp_path):
    from bank_mineru_mcp.parse_jobs import parse_job
    from bank_mineru_mcp.structured_store import StructuredStoreError
    task=tmp_path/'task_test';task.mkdir()
    path=task/'partial.xlsx';wb=Workbook();ws=wb.active
    ws.append(['id','amount','calculated']);ws.append([1,10,'=B2*2']);wb.save(path)
    source=SimpleNamespace(task_id='task_test',file_id='f',path=path,expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    store=StructuredStore(root=tmp_path)
    handle,inventory=await parse_job(store,source)
    assert inventory['sheets'][0]['formula_cache_status']=='partial'
    good=store.aggregate(handle.document_ref,[{'metrics':[{'column':'amount','fn':'sum'}]}])
    assert good['results'][0]['groups'][0]['amount:sum']==10
    with pytest.raises(StructuredStoreError) as error:
        store.aggregate(handle.document_ref,[{'metrics':[{'column':'calculated','fn':'sum'}]}])
    assert error.value.code=='DOCUMENT_FORMULA_CACHE_MISSING'
    assert store.read_range(handle.document_ref)['quality']=='partial'


def test_shared_string_table_is_read_from_disk(tmp_path):
    import zipfile
    path=tmp_path/'shared.xlsx';wb=Workbook();ws=wb.active
    ws.append(['id','text']);ws.append([1,'replace-me']);wb.save(path)
    with zipfile.ZipFile(path) as archive:
        parts={name:archive.read(name) for name in archive.namelist()}
    content=parts['[Content_Types].xml'].decode().replace('</Types>','<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    parts['[Content_Types].xml']=content.encode()
    parts['xl/sharedStrings.xml']=b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><t>shared </t></r><r><t>text</t></r></si></sst>'
    parts['xl/worksheets/sheet1.xml']=parts['xl/worksheets/sheet1.xml'].replace(b'<c r="B2" t="inlineStr"><is><t>replace-me</t></is></c>',b'<c r="B2" t="s"><v>0</v></c>')
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,data in parts.items():archive.writestr(name,data)
    inventory=extract_workbook(path,tmp_path/'work',stem='f')
    assert inventory['total_rows']==1
    row=json.loads((tmp_path/'work/sheet_00.rows.jsonl').read_text())
    assert row['v'][1]=='shared text'
    assert not list((tmp_path/'work').glob('.strings-*'))


@pytest.mark.asyncio
async def test_cancel_releases_worker_and_temporary_workspace(tmp_path):
    import asyncio
    from bank_mineru_mcp.parse_jobs import parse_job
    task=tmp_path/'task_test';task.mkdir();path=task/'source.csv'
    with path.open('w') as f:
        f.write('id,text\n')
        for i in range(40000):f.write(f'{i},'+('x'*1000)+'\n')
    source=SimpleNamespace(task_id='task_test',file_id='f',path=path,expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    running=asyncio.create_task(parse_job(StructuredStore(root=tmp_path),source))
    async with asyncio.timeout(10):
        while not any(json.loads(p.read_text()).get('status')=='running' for p in (task/'.reading-jobs').glob('*.json')):
            await asyncio.sleep(.01)
    running.cancel()
    with pytest.raises(asyncio.CancelledError):await running
    assert not list((task/'.reading-jobs').glob('*.work'))
    states=[json.loads(p.read_text()) for p in (task/'.reading-jobs').glob('*.json')]
    assert states[0]['status']=='cancelled'


@pytest.mark.asyncio
async def test_five_workbooks_keep_successful_references_with_large_directories(tmp_path):
    task = tmp_path / 'task_test'
    task.mkdir()
    sources = {}
    for index in range(5):
        path = task / f'file{index}.xlsx'
        wb = Workbook()
        wb.remove(wb.active)
        for sheet_index in range(40):
            ws = wb.create_sheet('建设规划工作表名称内容' * 2 + str(sheet_index))
            ws.append(['编号', '金额'])
            ws.append([1, 10])
        wb.save(path)
        sources[str(index)] = SimpleNamespace(task_id='task_test', file_id=str(index), path=path,
            extension='.xlsx', media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    service = MinerUToolService(file_resolver=SimpleNamespace(resolve=lambda ref: sources[ref]),
        mineru_client=None, document_store=DocumentStore(root=tmp_path),
        structured_store=StructuredStore(root=tmp_path))
    with approved_document_call('task_test', 'parse_documents', {}) as meta:
        with consume_document_call(meta, 'parse_documents', {}):
            result = await service.parse_documents([{'file_id': key, 'file_ref': key} for key in sources])
    assert result['status'] == 'completed'
    assert all(item['document_ref'] for item in result['items'])
    assert len(json.dumps(result, ensure_ascii=False, indent=2).encode()) < 32000


@pytest.mark.asyncio
async def test_merged_invalid_formula_remains_invalid_on_later_row(tmp_path):
    from bank_mineru_mcp.parse_jobs import parse_job
    from bank_mineru_mcp.structured_store import StructuredStoreError
    task = tmp_path / 'task_test'
    task.mkdir()
    path = task / 'merged-formula.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.append(['id', 'amount'])
    ws.append([1, '=1+2'])
    ws.append([2, None])
    ws.merge_cells('B2:B3')
    wb.save(path)
    source = SimpleNamespace(task_id='task_test', file_id='f', path=path,
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    store = StructuredStore(root=tmp_path)
    handle, _ = await parse_job(store, source)
    assert store.read_range(handle.document_ref, rows=[2,2])['quality'] == 'partial'
    with pytest.raises(StructuredStoreError) as caught:
        store.aggregate(handle.document_ref, [{'filter': {'column':'id','op':'eq','value':2},
            'metrics':[{'column':'amount','fn':'sum'}]}])
    assert caught.value.code == 'DOCUMENT_FORMULA_CACHE_MISSING'


@pytest.mark.asyncio
async def test_concurrent_retries_publish_only_one_document(tmp_path):
    import asyncio
    from bank_mineru_mcp.parse_jobs import parse_job
    task = tmp_path / 'task_test'
    task.mkdir()
    path = task / 'source.csv'
    path.write_text('id,amount\n1,10\n')
    source = SimpleNamespace(task_id='task_test', file_id='f', path=path,
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    store = StructuredStore(root=tmp_path)
    results = await asyncio.gather(parse_job(store,source), parse_job(store,source))
    assert results[0][0].document_ref == results[1][0].document_ref
    assert len(list((task/'.mineru-struct').glob('*/manifest.json'))) == 1


def test_high_cardinality_groups_are_losslessly_paged(tmp_path):
    task=tmp_path/'task_test'
    task.mkdir()
    path=task/'groups.csv'
    path.write_text('id,amount\n'+''.join(f'{i},1\n' for i in range(35)))
    inventory=extract_workbook(path,task/'work',stem='f')
    source=SimpleNamespace(task_id='task_test',path=path,expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    store=StructuredStore(root=tmp_path,max_groups=10)
    handle=store.write(source,inventory,task/'work')
    offset=0
    groups=[]
    while True:
        response=store.aggregate(handle.document_ref,[{'group_by':['id'],
            'metrics':[{'column':'amount','fn':'sum'}],'group_cursor':offset}])
        result=response['results'][0]
        groups.extend(result['groups'])
        if result['next_group_cursor'] is None:
            break
        assert result['next_group_cursor']>offset
        offset=result['next_group_cursor']
    assert len(groups)==35
    assert {str(item['group']['id']) for item in groups}=={str(i) for i in range(35)}
    assert sum(item['amount:sum'] for item in groups)==35
