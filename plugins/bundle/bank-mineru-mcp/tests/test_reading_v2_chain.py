"""Real MCP wire -> native Driver adapter -> read ledger regression."""
import json
import socket
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from mcp import ClientSession
from mcp.client.streamable_http import streamablehttp_client
from openpyxl import Workbook

from bank_mineru_mcp.config import MinerUSettings
from bank_mineru_mcp.document_store import DocumentStore
from bank_mineru_mcp.structured_store import StructuredStore
from bank_mineru_mcp.tools import MinerUToolService
from bank_mineru_mcp.server import MinerUMcpService
from bank_runtime.gateway.document_reads import DocumentReadLedger
from qwenpaw.drivers.adapters.agentscope_tool import _blocks_from_value


class NoLayoutClient:
    async def parse(self, *args, **kwargs):
        raise AssertionError("Excel must never call the MinerU layout engine")

    async def probe(self):
        return {"status": "healthy"}

    async def close(self):
        pass


@pytest.mark.asyncio
@pytest.mark.parametrize('same_task', [True, False])
async def test_reading5_same_name_revision_statistics_and_raw_coverage(tmp_path, same_task):
    from bank_runtime.gateway.document_access import approved_document_call, consume_document_call
    from bank_mineru_mcp.tools import ToolContractError

    sources = {}
    for revision in (0, 1):
        task_id = 'task_original' if same_task or revision == 0 else 'task_revised'
        task = tmp_path / task_id
        task.mkdir(exist_ok=True)
        path = task / f'file_{revision}.xlsx'
        workbook = Workbook(); sheet = workbook.active; sheet.title = '台账'
        sheet.append(['记录编号', '问题大类', '问题小类', '问题描述', '风险等级', '机构', '发现日期', '金额（元）'])
        for index in range(1, 13):
            risk = ['高', '中', '低'][(index - 1) % 3]
            amount = index * 100
            if revision and index == 1:
                risk, amount = '低', 1100
            sheet.append([f'T{index:03d}', '测试大类', '测试小类', f'记录{index}', risk,
                          '测试机构', '2026-01-01', amount])
        workbook.save(path)
        sources[f'authorized-{revision}'] = SimpleNamespace(
            task_id=task_id, file_id=f'file_{revision}', path=path,
            original_name='中文问题台账（12行）.xlsx', extension='.xlsx',
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    service = MinerUToolService(file_resolver=SimpleNamespace(resolve=sources.__getitem__),
        mineru_client=NoLayoutClient(), document_store=DocumentStore(root=tmp_path),
        structured_store=StructuredStore(root=tmp_path))
    ledgers = [DocumentReadLedger(), DocumentReadLedger()]

    async def invoke(revision, name, arguments):
        source = sources[f'authorized-{revision}']
        with approved_document_call(source.task_id, name, arguments) as grant:
            with consume_document_call(grant, name, arguments):
                value = (await service.parse_documents(**arguments) if name == 'parse_documents'
                         else await service.execute_structured_query(name, arguments))
        ledger = ledgers[revision]
        ledger.start('MinerU__' + name, arguments)
        ledger.observe('MinerU__' + name, arguments, [{'type': 'text', 'text': json.dumps(value)}], True)
        return value

    refs, statistics = [], []
    for revision, expected in [(0, {'高': 4, '中': 4, '低': 4}), (1, {'高': 3, '中': 4, '低': 5})]:
        item = (await invoke(revision, 'parse_documents', {'documents': [
            {'file_id': f'file_{revision}', 'file_ref': f'authorized-{revision}'}]}))['items'][0]
        assert item['status'] == 'completed'
        assert item['inventory']['title'] == '中文问题台账（12行）.xlsx'
        ref = item['document_ref']; refs.append(ref)
        args = {'document_ref': ref, 'ops': [{'sheet': '台账', 'group_by': ['风险等级'],
            'metrics': [{'column': '记录编号', 'fn': 'count'}, {'column': '金额（元）', 'fn': 'sum'}, {'column': '金额（元）', 'fn': 'avg'}]}]}
        stats = await invoke(revision, 'aggregate', args); statistics.append((args, stats))
        groups = stats['results'][0]['groups']
        assert {g['group']['风险等级']: g['记录编号:count'] for g in groups} == expected
        assert sum(g['金额（元）:sum'] for g in groups) == (8800 if revision else 7800)
        ledger = ledgers[revision]
        assert not ledger.pending and not ledger.documents[ref].complete
        table = '| 风险等级 | 记录数 | 金额合计（元） | 平均金额（元） |\n| --- | --- | --- | --- |\n'
        table += '\n'.join(f"| {g['group']['风险等级']} | {g['记录编号:count']} | {g['金额（元）:sum']} | {g['金额（元）:avg']} |" for g in groups)
        assert ledger.declaration_conflict(table) == ''
        assert ledger.declaration_conflict('已完整读取台账全文。' + table) == 'DOCUMENT_READ_INCOMPLETE'
        page = await invoke(revision, 'read_range', {'document_ref': ref, 'sheet': '台账', 'rows': [1, 12]})
        assert page['all_columns'] is True and page['rows_returned'] == [1, 12]
        assert ledger.documents[ref].complete
        assert ledger.declaration_conflict('已完整读取台账全文。\n' + table) == ''
    assert refs[0] != refs[1]
    # Replaying the original query after the revision must retain its original result.
    assert await invoke(0, 'aggregate', statistics[0][0]) == statistics[0][1]
    assert await invoke(1, 'aggregate', statistics[1][0]) == statistics[1][1]
    if not same_task:
        with pytest.raises(ToolContractError) as denied:
            await invoke(1, 'aggregate', statistics[0][0])
        assert denied.value.code == 'FILE_ACCESS_DENIED'


@pytest.mark.asyncio
async def test_five_sheet_excel_wire_pagination_and_coverage(tmp_path):
    task = tmp_path / "task_test"
    task.mkdir()
    path = task / "planning.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    expected = {}
    for name, count in zip(("规划总览", "年度规划", "节奏矩阵", "技术底座", "事项清单"), (17, 65, 14, 14, 15)):
        sheet = workbook.create_sheet(name)
        sheet.append(["编号", "说明"])
        for index in range(count):
            sheet.append([index + 1, "技术建设内容" * 80])
        expected[name] = count
    workbook.save(path)
    source = SimpleNamespace(task_id="task_test", file_id="file_test", path=path,
        extension=".xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    client = NoLayoutClient()
    service = MinerUToolService(file_resolver=SimpleNamespace(resolve=lambda ref: source),
        mineru_client=client, document_store=DocumentStore(root=tmp_path),
        structured_store=StructuredStore(root=tmp_path))
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    server = MinerUMcpService(settings=MinerUSettings(base_url="http://unused.test", submit_mode="file_parse", token="test", mcp_port=port),
        tool_service=service, mineru_client=client)
    await server.start()
    try:
        async with streamablehttp_client(f"http://127.0.0.1:{port}/mcp") as (read, write, _):
            async with ClientSession(read, write) as session:
                await session.initialize()
                ledger = DocumentReadLedger()
                async def call(name, payload):
                    ledger.start("MinerU__" + name, payload)
                    from bank_runtime.gateway.document_access import approved_document_call
                    with approved_document_call("task_test", name, payload) as metadata:
                        result = await session.call_tool(name, payload, meta=metadata)
                    blocks = _blocks_from_value(result)
                    assert len(blocks) == 1
                    assert len(blocks[0].text.encode()) <= 32000
                    ledger.observe("MinerU__" + name, payload, blocks, not result.isError)
                    return json.loads(blocks[0].text)
                parsed = await call("parse_documents", {"documents": [{"file_id": "file_test", "file_ref": "authorized-test-ref"}]})
                item = parsed["items"][0]
                assert item["content_mode"] == "structured"
                assert item["inventory"]["engine"] == "ooxml-1"
                ref = item["document_ref"]
                invalid = await call("read_range", {"document_ref": ref, "sheet": "规划总览", "columns": ["编号", "编号"]})
                assert invalid["error_code"] == "DOCUMENT_ARGUMENT_INVALID"
                assert not ledger.documents[ref].complete
                projected = await call("read_range", {"document_ref": ref, "sheet": "规划总览", "columns": ["编号"]})
                assert projected["all_columns"] is False
                assert ledger.documents[ref].covered_rows("规划总览") == 0
                assert ledger.declaration_conflict("所有工作表全量总结") == "DOCUMENT_READ_INCOMPLETE"
                cursor = None
                pages = 0
                while True:
                    payload = {"document_ref": ref, "cursor": cursor, "limit": 10}
                    page = await call("read_document_chunks", payload)
                    assert page == await call("read_document_chunks", payload)
                    pages += 1
                    if not page["has_more"]:
                        break
                    cursor = page["next_cursor"]
                assert pages > 1
                assert not ledger.pending
                assert ledger.documents[ref].complete
                assert {name: ledger.documents[ref].covered_rows(name) for name in expected} == expected
                assert ledger.declaration_conflict("所有工作表全量总结") == ""
                # A restarted store can resume the same authorized reference.
                restarted = StructuredStore(root=tmp_path)
                assert restarted.read_range(ref, sheet="年度规划", rows=[1, 3])["rows_scanned"] == 3
    finally:
        await server.stop()


@pytest.mark.asyncio
async def test_merged_multirow_workbook_short_refs_and_metric_alias_through_gateway(tmp_path, monkeypatch):
    from copy import deepcopy
    from agentscope.message import ToolCallBlock, ToolResultState
    from agentscope.permission import PermissionBehavior, PermissionDecision
    from agentscope.tool import ToolResponse
    from bank_runtime.gateway.middleware import BankRuntimeGatewayMiddleware, GatewayPermissionEngine
    from bank_runtime.sandbox import file_refs
    from bank_runtime.sandbox.file_refs import FileRefRegistry
    from bank_runtime.sandbox.cache import PreparedSandboxFile
    from qwenpaw.drivers.mcp_context import current_mcp_metadata

    root=tmp_path/'task_report'; root.mkdir(); path=root/'report.xlsx'
    book=Workbook(); book.remove(book.active)
    for name, values in [('本期',[10,20]),('上期',[3,4])]:
        sheet=book.create_sheet(name); sheet.append(['经营报表',None]); sheet.merge_cells('A1:B1')
        sheet.append(['部门','金额'])
        for i,v in enumerate(values): sheet.append([f'部门{i}',v])
    book.save(path)
    registry=FileRefRegistry(root=tmp_path); monkeypatch.setattr(file_refs,'_REGISTRY',registry)
    registry.issue(PreparedSandboxFile(file_id='file_report',local_path=path,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        size_bytes=path.stat().st_size,original_name=path.name,expires_at='',task_id='task_report'),
        expires_at=datetime.now(timezone.utc)+timedelta(minutes=10))
    parser=NoLayoutClient()
    service=MinerUToolService(file_resolver=registry,mineru_client=parser,document_store=DocumentStore(root=tmp_path),
        structured_store=StructuredStore(root=tmp_path))
    with socket.socket() as sock:
        sock.bind(('127.0.0.1',0)); port=sock.getsockname()[1]
    server=MinerUMcpService(settings=MinerUSettings(base_url='http://unused.test',submit_mode='file_parse',token='test',mcp_port=port),
        tool_service=service,mineru_client=parser)
    class Gateway:
        config=SimpleNamespace(task_id='task_report')
        admitted=None
        async def preflight(self,name,args,**kw): self.admitted=deepcopy(args); return {'tool_call_id':'call'}
        async def report_guard(self,*args): pass
        async def report_result(self,*args): pass
    class Guard:
        async def check_permission(self,*args): return PermissionDecision(behavior=PermissionBehavior.ALLOW,message='allow')
    gateway=Gateway(); middleware=BankRuntimeGatewayMiddleware(gateway); engine=GatewayPermissionEngine(Guard(),middleware)
    await server.start()
    try:
        async with streamablehttp_client(f'http://127.0.0.1:{port}/mcp') as (read,write,_):
            async with ClientSession(read,write) as session:
                await session.initialize()
                async def invoke(name,args):
                    tool='MinerU__'+name
                    assert (await engine.check_permission(SimpleNamespace(name=tool),args)).behavior==PermissionBehavior.ALLOW
                    async def handler(**kw):
                        actual=json.loads(kw['tool_call'].input) if kw else args
                        assert actual==gateway.admitted
                        result=await session.call_tool(name,actual,meta=current_mcp_metadata())
                        yield ToolResponse(id='call',state=ToolResultState.SUCCESS,content=_blocks_from_value(result))
                    output=[item async for item in middleware.on_acting(None,{'tool_call':ToolCallBlock(id='call',name=tool,input=json.dumps(args))},handler)]
                    return json.loads(output[0].content[0].text)
                result=await invoke('parse_documents',{'documents':[{'file_id':'file_report'}]})
                item=result['items'][0]; assert item['status']=='completed'
                for meta,expected in zip(item['inventory']['sheets'],[30,7]):
                    name=meta['name']; column=meta['columns'][1]['name']
                    first=await invoke('read_range',{'document_ref':'file_report','sheet':name,'rows':['1','1']})
                    assert first.get('status')!='failed'
                    stats=await invoke('aggregate',{'document_ref':'file_report','ops':[{'sheet':name,'row_range':[1,meta['rows']],
                        'metrics':[{'column':column,'op':'sum'}]}]})
                    assert stats['results'][0]['groups'][0][column+':sum']==expected
                    assert not middleware.document_reads.pending
                # Scoped statistics cannot claim all raw cells were read.
                assert not middleware.document_reads.documents[item['document_ref']].complete
    finally:
        await server.stop()


@pytest.mark.asyncio
@pytest.mark.parametrize('row_count,column_count', [(7445, 4), (600, 13)])
async def test_ledger_workbook_batch_statistics_finish_without_reading_every_row(tmp_path, monkeypatch, row_count, column_count):
    from unittest.mock import AsyncMock
    from bank_mineru_mcp import parse_jobs
    from bank_runtime.gateway.document_access import approved_document_call, consume_document_call
    root = tmp_path / 'task_statistics'; root.mkdir()
    path = root / 'file_opaque.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = 'Sheet1'
    sheet.append(['问题大类', '问题小类', '问题描述', '风险等级'] + [f'补充字段{i}' for i in range(column_count-4)])
    for index in range(row_count):
        sheet.append([f'大类{index%3}', f'小类{index%7}', f'记录{index}', f'风险{index%2}'] + [index for _ in range(column_count-4)])
    book.create_sheet('Sheet2'); book.create_sheet('Sheet3'); book.save(path)
    source = SimpleNamespace(task_id='task_statistics', file_id='file_statistics', path=path,
        original_name='总行问题台账历史流水.xlsx', extension='.xlsx',
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    service = MinerUToolService(file_resolver=SimpleNamespace(resolve=lambda ref: source),
        mineru_client=NoLayoutClient(), document_store=DocumentStore(root=tmp_path), structured_store=StructuredStore(root=tmp_path))
    ledger = DocumentReadLedger()
    spy = AsyncMock(wraps=parse_jobs.query_job); monkeypatch.setattr(parse_jobs, 'query_job', spy)
    async def invoke(name, arguments):
        with approved_document_call(source.task_id, name, arguments) as authorization:
            with consume_document_call(authorization, name, arguments):
                value = await service.parse_documents(**arguments) if name == 'parse_documents' else await service.execute_structured_query(name, arguments)
        ledger.start('MinerU__' + name, arguments)
        ledger.observe('MinerU__' + name, arguments, [{'type':'text', 'text':json.dumps(value, ensure_ascii=False)}], True)
        return value
    parsed = await invoke('parse_documents', {'documents':[{'file_id':source.file_id, 'file_ref':'authorized-file'}]})
    item = parsed['items'][0]; ref = item['document_ref']
    assert item['inventory']['title'] == source.original_name
    assert item['inventory']['total_rows'] == row_count
    metrics = [{'column':'问题描述', 'fn':'count'}]
    operations = [{'sheet':'Sheet1', 'group_by':group, 'metrics':metrics}
                  for group in [['问题大类','问题小类'], ['风险等级'], []]]
    result = await invoke('aggregate', {'document_ref':ref, 'ops':operations})
    assert len(result['results']) == 3 and result['truncated'] is False
    for operation in result['results']:
        assert sum(group['问题描述:count'] for group in operation['groups']) == row_count
    assert not ledger.pending
    assert ledger.declaration_conflict(f'全表按问题大类/问题小类统计问题描述数量合计{row_count}条') == ''
    assert ledger.declaration_conflict('全表按风险等级统计问题描述计数') == ''
    assert ledger.declaration_conflict(f'全表合计{row_count:,}条') == ''
    assert ledger.declaration_conflict('已完整读取全部内容') == 'DOCUMENT_READ_INCOMPLETE'
    assert ledger.documents[ref].covered_rows('Sheet1') == 0
    assert spy.await_count == 1  # One batch; no find/read/statistics loop.
    assert await invoke('aggregate', {'document_ref':ref, 'ops':operations}) == result
    assert spy.await_count == 1  # A retry reuses results, without another scan.
