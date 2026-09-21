"""Structured store regression: extraction, bounded reads, aggregates, recovery."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest

from bank_mineru_mcp.spreadsheet import extract_workbook
from bank_mineru_mcp.structured_store import StructuredStore, StructuredStoreError

openpyxl = pytest.importorskip("openpyxl")


def _source(tmp_path: Path, name: str, suffix: str, mime: str):
    task_root = tmp_path / "task_001"
    task_root.mkdir(parents=True, exist_ok=True)
    path = task_root / name
    return SimpleNamespace(
        task_id="task_001",
        file_id="file_001",
        path=path,
        media_type=mime,
        extension=suffix,
        size_bytes=0,
        sha256="0" * 64,
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1),
    )


def _write_csv(path: Path, rows: int = 60) -> None:
    lines = ["姓名,团队,营销笔数,营销金额"]
    for index in range(1, rows + 1):
        lines.append(f"员工{index:03d},团队{(index - 1) // 10 + 1},{index},{index * 10.5}")
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("支行01", "支行02"):
        sheet = workbook.create_sheet(name)
        sheet.merge_cells("A1:D1")
        sheet["A1"] = f"{name}营销数据"
        for column, title in enumerate(["姓名", "岗位", "营销笔数", "营销金额"], start=1):
            sheet.cell(row=2, column=column, value=title)
        for row in range(3, 23):
            if (row - 3) % 10 == 0:
                sheet.merge_cells(start_row=row, start_column=2, end_row=row + 9, end_column=2)
                sheet.cell(row=row, column=2, value="团队1" if row < 13 else "团队2")
            sheet.cell(row=row, column=1, value=f"员工{row:03d}")
            sheet.cell(row=row, column=3, value=row)
            sheet.cell(row=row, column=4, value=row * 2.0)
    workbook.save(path)


def _store(tmp_path: Path, **kwargs):
    return StructuredStore(root=tmp_path, **kwargs)


def _parse_csv(tmp_path: Path):
    source = _source(tmp_path, "book.csv", ".csv", "text/csv")
    _write_csv(source.path)
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path)
    handle = store.write(source, inventory, work)
    return store, handle


def test_csv_inventory_and_range_paging(tmp_path) -> None:
    store, handle = _parse_csv(tmp_path)
    inventory = store.inventory(handle.document_ref)
    assert inventory["sheet_count"] == 1
    assert inventory["sheets"][0]["rows"] == 60
    assert inventory["sheets"][0]["header_row"] == 1
    first = store.read_range(handle.document_ref, rows=[1, 40])
    assert first["rows_returned"] == [1, 40]
    assert first["markdown"].count("|") > 10
    assert "姓名" in first["markdown"]
    second = store.read_range(handle.document_ref, rows=[41, 60])
    assert second["rows_returned"] == [41, 60]
    assert second["has_more"] is False
    assert second["sheet_total_rows"] == 60
    paged = store.read_range(handle.document_ref, row_cursor=11)
    assert paged["rows_returned"][0] == 11


def test_range_page_respects_char_budget(tmp_path) -> None:
    source = _source(tmp_path, "big.csv", ".csv", "text/csv")
    _write_csv(source.path, rows=4000)
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path, page_chars=4000)
    handle = store.write(source, inventory, work)
    page = store.read_range(handle.document_ref)
    assert len(page["markdown"]) <= 4100
    assert page["has_more"] is True


def test_xlsx_merged_ranges_expand(tmp_path) -> None:
    source = _source(
        tmp_path,
        "book.xlsx",
        ".xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    _write_xlsx(source.path)
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path)
    handle = store.write(source, inventory, work)
    sheet = store.inventory(handle.document_ref)["sheets"][0]
    assert sheet["merged_ranges"], "merged ranges must be preserved"
    page = store.read_range(handle.document_ref, sheet="支行01", rows=[1, 20])
    assert page["markdown"].count("团队1") == 10, "row-span values must fill down"
    assert page["sheet"] == "支行01"


def test_aggregate_metrics_median_filter_and_union(tmp_path) -> None:
    source = _source(
        tmp_path,
        "book.xlsx",
        ".xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    _write_xlsx(source.path)
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path)
    handle = store.write(source, inventory, work)
    result = store.aggregate(
        handle.document_ref,
        [
            {
                "sheet": "支行01",
                "group_by": ["岗位"],
                "metrics": [
                    {"column": "营销笔数", "fn": "sum"},
                    {"column": "营销金额", "fn": "median"},
                    {"column": "姓名", "fn": "count_distinct"},
                ],
                "filter": {"column": "营销笔数", "op": "gte", "value": 3},
            },
            {
                "sheet": "*",
                "cross_sheet_union": {"key_column": "姓名"},
                "metrics": [{"column": "营销笔数", "fn": "count"}],
            },
        ],
    )
    grouped, union = result["results"]
    assert grouped["rows_matched"] == 20
    assert grouped["full_range"] is True
    assert union["rows_scanned"] == 40
    assert union["rows_matched"] == 40
    first_group = grouped["groups"][0]
    assert first_group["group"]["岗位"] == "团队1"
    assert first_group["营销笔数:sum"] == sum(range(3, 13))
    assert first_group["营销金额:median"] == 15.0
    assert first_group["姓名:count_distinct"] == 10


def test_aggregate_invalid_metric_rejected(tmp_path) -> None:
    store, handle = _parse_csv(tmp_path)
    with pytest.raises(StructuredStoreError) as failed:
        store.aggregate(handle.document_ref, [{"sheet": "data", "metrics": [{"column": "营销笔数", "fn": "stddev"}]}])
    assert failed.value.code == "DOCUMENT_ARGUMENT_INVALID"


def test_search_locates_rows(tmp_path) -> None:
    store, handle = _parse_csv(tmp_path)
    hits = store.search(handle.document_ref, query="员工042")
    assert len(hits["hits"]) == 1
    assert hits["hits"][0]["row"] == 42


def test_read_chunks_shim_covers_blocks(tmp_path) -> None:
    store, handle = _parse_csv(tmp_path)
    seen = []
    cursor = None
    while True:
        page = store.read_chunks(handle.document_ref, cursor=cursor, limit=10)
        seen.extend(chunk["index"] for chunk in page["chunks"])
        cursor = page["next_cursor"]
        if not page["has_more"]:
            break
    assert seen == list(range(len(seen)))
    assert page["coverage"]["total"] == len(seen)


def test_registry_survives_restart_and_cursor_tamper_rejected(tmp_path) -> None:
    source = _source(tmp_path, "book.csv", ".csv", "text/csv")
    _write_csv(source.path, rows=200)
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path, page_chars=1000)
    handle = store.write(source, inventory, work)
    page = store.read_range(handle.document_ref)
    assert page["has_more"] is True
    restarted = _store(tmp_path, page_chars=1000)
    again = restarted.read_range(handle.document_ref, row_cursor=page["next_row_cursor"])
    assert again["rows_returned"][0] == page["next_row_cursor"]
    with pytest.raises(StructuredStoreError) as failed:
        restarted.read_chunks(
            handle.document_ref, cursor="cs1_5_" + "0" * 32 + "_" + "0" * 64, limit=1
        )
    assert failed.value.code == "FILE_REF_INVALID"


def test_expired_reference_rejected(tmp_path) -> None:
    store, handle = _parse_csv(tmp_path)
    now = datetime.now(timezone.utc)
    store.clock = lambda: now + timedelta(hours=2)
    with pytest.raises(StructuredStoreError) as failed:
        store.read_range(handle.document_ref)
    assert failed.value.code == "DOCUMENT_REF_EXPIRED"


def test_legacy_unicode_pages_are_bounded_and_complete(tmp_path):
    import json
    source = _source(tmp_path, "wide.csv", ".csv", "text/csv")
    source.path.write_text("名称,说明\n" + "\n".join(f"条目{i}," + "中文内容" * 150 for i in range(80)), encoding="utf-8")
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path)
    handle = store.write(source, inventory, work)
    cursor = None
    seen = []
    while True:
        page = store.read_chunks(handle.document_ref, cursor=cursor, limit=10)
        assert len(json.dumps(page, ensure_ascii=False, indent=2).encode()) <= 32000
        assert page == store.read_chunks(handle.document_ref, cursor=cursor, limit=10)
        for chunk in page["chunks"]:
            start, end = chunk["rows_returned"]
            seen.extend(range(start, end + 1))
        if not page["has_more"]:
            break
        assert page["next_cursor"] != cursor
        cursor = page["next_cursor"]
    assert seen == list(range(1, 81))


def test_range_unicode_response_stays_bounded(tmp_path):
    import json
    source = _source(tmp_path, "wide.csv", ".csv", "text/csv")
    source.path.write_text("名称,说明\n" + "\n".join(f"条目{i}," + "中文内容" * 150 for i in range(80)), encoding="utf-8")
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path)
    ref = store.write(source, inventory, work).document_ref
    page = store.read_range(ref)
    assert len(json.dumps(page, ensure_ascii=False, indent=2).encode()) <= 32000
    assert page["has_more"]


def test_blank_rows_keep_stable_coordinates_and_long_cells(tmp_path):
    source = _source(tmp_path, "blank.csv", ".csv", "text/csv")
    source.path.write_text("编号,说明\n1," + "长" * 3000 + "\n,\n3,结尾", encoding="utf-8")
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path)
    ref = store.write(source, inventory, work).document_ref
    result = store.read_range(ref, format="records")
    assert result["sheet_total_rows"] == 3
    assert [row["row"] for row in result["records"]] == [1, 2, 3]
    assert len(result["records"][0]["values"][1]) == 3000


def test_single_oversized_row_returns_error_without_truncation(tmp_path):
    source = _source(tmp_path, "largecell.csv", ".csv", "text/csv")
    source.path.write_text("编号,说明\n1," + "长" * 15000, encoding="utf-8")
    work = tmp_path / "work"
    inventory = extract_workbook(source.path, work, stem="file_001")
    store = _store(tmp_path)
    ref = store.write(source, inventory, work).document_ref
    with pytest.raises(StructuredStoreError) as caught:
        store.read_range(ref)
    assert caught.value.code == "DOCUMENT_RESULT_TOO_LARGE"
    assert store.read_range(ref, columns=["编号"])["has_more"] is False

@pytest.mark.parametrize('columns', [['姓名', '姓名', '姓名', '姓名'], ['姓名', '不存在']])
def test_projection_rejects_duplicate_or_unknown_columns(tmp_path, columns):
    store, handle = _parse_csv(tmp_path)
    with pytest.raises(StructuredStoreError, match='columns') as error:
        store.read_range(handle.document_ref, columns=columns)
    assert error.value.code == 'DOCUMENT_ARGUMENT_INVALID'

@pytest.mark.parametrize('op', [
    {'metrics': [{'column': '不存在', 'fn': 'sum'}]},
    {'group_by': ['不存在'], 'metrics': [{'column': '营销金额', 'fn': 'sum'}]},
    {'row_range': [2, 1], 'metrics': [{'column': '营销金额', 'fn': 'sum'}]},
    {'metrics': 'sum'}, {'metrics': []}, {'group_by': '姓名'},
    {'filter': {'column': '姓名', 'op': 'in', 'value': 3}, 'metrics': [{'column': '姓名', 'fn': 'count'}]},
])
def test_invalid_aggregate_is_recoverable_contract_error(tmp_path, op):
    store, handle = _parse_csv(tmp_path)
    with pytest.raises(StructuredStoreError) as error:
        store.aggregate(handle.document_ref, [op])
    assert error.value.code == 'DOCUMENT_ARGUMENT_INVALID'
    assert store.aggregate(handle.document_ref, [{'metrics': [{'column': '营销笔数', 'fn': 'sum'}]}])['results'][0]['groups'][0]['营销笔数:sum'] == 1830


def test_missing_formula_cache_cannot_become_successful_partial_sum(tmp_path):
    from bank_mineru_mcp.spreadsheet import SpreadsheetExtractError
    source = _source(tmp_path, 'formula.xlsx', '.xlsx', '')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['姓名', '金额'])
    ws.append(['甲', '=1+2'])
    ws.append(['乙', 5])
    wb.save(source.path)
    with pytest.raises(SpreadsheetExtractError) as error:
        extract_workbook(source.path, tmp_path / 'work', stem='formula')
    assert error.value.code == 'DOCUMENT_FORMULA_CACHE_MISSING'


def test_structured_quota_rejects_and_removes_temporary_result(tmp_path):
    source = _source(tmp_path, 'book.csv', '.csv', 'text/csv')
    _write_csv(source.path, rows=300)
    work = tmp_path / 'work'
    inventory = extract_workbook(source.path, work, stem='one')
    store = _store(tmp_path, max_task_bytes=1024)
    with pytest.raises(StructuredStoreError) as error:
        store.write(source, inventory, work)
    assert error.value.code == 'DOCUMENT_RESULT_TOO_LARGE'
    assert not work.exists()
    assert not list((tmp_path / 'task_001' / '.mineru-struct').glob('*/manifest.json'))


def test_structured_quota_counts_existing_files_after_restart(tmp_path):
    source = _source(tmp_path, 'book.csv', '.csv', 'text/csv')
    _write_csv(source.path, rows=60)
    work = tmp_path / 'one'
    inventory = extract_workbook(source.path, work, stem='one')
    store = _store(tmp_path)
    handle = store.write(source, inventory, work)
    used = sum(p.stat().st_size for p in handle.path.rglob('*') if p.is_file())
    store = _store(tmp_path, max_task_bytes=used + 100)
    work = tmp_path / 'two'
    inventory = extract_workbook(source.path, work, stem='two')
    with pytest.raises(StructuredStoreError):
        store.write(source, inventory, work)
    assert store.read_range(handle.document_ref)['rows_scanned'] == 60
    assert not work.exists()

@pytest.mark.parametrize('cached', [True, False])
def test_formula_cache_and_real_blank_remain_distinct(tmp_path, cached):
    import zipfile
    from bank_mineru_mcp.spreadsheet import SpreadsheetExtractError
    source = _source(tmp_path, 'cached.xlsx', '.xlsx', '')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['姓名', '金额'])
    ws.append(['甲', '=1+2'])
    ws.append(['乙', 5])
    ws.append(['丙', None])
    wb.save(source.path)
    if cached:
        with zipfile.ZipFile(source.path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        from xml.etree import ElementTree as ET
        xml = ET.fromstring(members['xl/worksheets/sheet1.xml'])
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        xml.find('.//s:c[@r="B2"]/s:v', ns).text = '3'
        members['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
        with zipfile.ZipFile(source.path, 'w') as archive:
            for name, content in members.items():
                archive.writestr(name, content)
    work = tmp_path / 'work'
    if not cached:
        with pytest.raises(SpreadsheetExtractError):
            extract_workbook(source.path, work, stem='cached')
        return
    inventory = extract_workbook(source.path, work, stem='cached')
    assert inventory['sheets'][0]['formula_count'] == 1
    assert inventory['sheets'][0]['formula_cache_status'] == 'available'
    store = _store(tmp_path)
    handle = store.write(source, inventory, work)
    page = store.read_range(handle.document_ref, format='records')
    assert [r['values'][1] for r in page['records']] == [3, 5, None]
    result = store.aggregate(handle.document_ref, [{'metrics': [{'column': '金额', 'fn': 'sum'}]}])
    assert result['results'][0]['groups'][0]['金额:sum'] == 8


def test_quota_serializes_competing_writers(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    source = _source(tmp_path, 'book.csv', '.csv', 'text/csv')
    _write_csv(source.path, rows=60)
    work = tmp_path / 'baseline'
    inventory = extract_workbook(source.path, work, stem='baseline')
    baseline = _store(tmp_path)
    handle = baseline.write(source, inventory, work)
    size = sum(p.stat().st_size for p in handle.path.rglob('*') if p.is_file())
    baseline.delete_task(source.task_id)
    # Each candidate fits alone; both together exceed the same task quota.
    stores = [_store(tmp_path, max_task_bytes=size + 100) for _ in range(2)]
    def write(index):
        work = tmp_path / f'work-{index}'
        inventory = extract_workbook(source.path, work, stem='file')
        try:
            stores[index].write(source, inventory, work)
            return 'ok'
        except StructuredStoreError as exc:
            assert not work.exists()
            return exc.code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(write, range(2))) == ['DOCUMENT_RESULT_TOO_LARGE', 'ok']


def test_formula_unaware_cached_workbook_requires_reparse(tmp_path):
    source = _source(tmp_path, 'book.xlsx', '.xlsx', '')
    _write_xlsx(source.path)
    work = tmp_path / 'old'
    inventory = extract_workbook(source.path, work, stem='old')
    for sheet in inventory['sheets']:
        sheet.pop('formula_cache_status')
        sheet.pop('formula_count')
    store = _store(tmp_path)
    old = store.write(source, inventory, work)
    with pytest.raises(StructuredStoreError) as error:
        store.read_range(old.document_ref, sheet='支行01')
    assert error.value.code == 'DOCUMENT_REF_EXPIRED'
    work = tmp_path / 'current'
    inventory = extract_workbook(source.path, work, stem='current')
    current = store.write(source, inventory, work)
    assert store.read_range(current.document_ref, sheet='支行01')['all_columns']


def test_header_aliases_remain_unique_and_address_the_same_physical_column(tmp_path):
    source = _source(tmp_path, 'headers.csv', '.csv', 'text/csv')
    source.path.write_text('金额,金额,金额#2,,col4\n10,20,30,40,50\n', encoding='utf-8')
    work = tmp_path / 'work'
    inventory = extract_workbook(source.path, work, stem='f')
    store = _store(tmp_path)
    handle = store.write(source, inventory, work)
    columns = inventory['sheets'][0]['columns']
    names = [c['name'] for c in columns]
    assert len(set(names)) == 5
    for column, expected in zip(columns, [10,20,30,40,50]):
        page = store.read_range(handle.document_ref, columns=[column['name']], format='records')
        result = store.aggregate(handle.document_ref, [{'metrics':[{'column':column['name'],'fn':'sum'}]}])
        assert page['records'][0]['values'] == [str(expected)]
        assert result['results'][0]['groups'][0][column['name']+':sum'] == expected
    assert [c['source_index'] for c in columns] == list(range(5))
    assert [c['original_name'] for c in columns] == ['金额','金额','金额#2','','col4']


@pytest.mark.parametrize('suffix,separator', [('.csv',','),('.tsv','\t')])
@pytest.mark.parametrize('newline', ['\n','\r\n'])
def test_delimited_quoted_cells_preserve_linebreaks_quotes_and_separators(tmp_path, suffix, separator, newline):
    import csv, io
    source = _source(tmp_path, 'lines'+suffix, suffix, 'text/plain')
    value = '第一行'+newline+'第二行'+separator+'"原文"'
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, delimiter=separator, lineterminator=newline)
    writer.writerows([['姓名','备注'], ['甲',value]])
    source.path.write_bytes(stream.getvalue().encode('utf-8'))
    work = tmp_path/'work'; inventory = extract_workbook(source.path, work, stem='f')
    store = _store(tmp_path); handle = store.write(source, inventory, work)
    assert store.read_range(handle.document_ref, format='records')['records'][0]['values'][1] == value
    assert len(store.search(handle.document_ref, query='第一行'+newline+'第二行')['hits']) == 1


def test_union_echoes_every_source_sheet_instead_of_first_sheet(tmp_path):
    source=_source(tmp_path,'book.xlsx','.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    _write_xlsx(source.path)
    work=tmp_path/'work'; inventory=extract_workbook(source.path,work,stem='f')
    store=_store(tmp_path); handle=store.write(source,inventory,work)
    result=store.aggregate(handle.document_ref,[{'cross_sheet_union':{'key_column':'姓名'},'metrics':[{'column':'营销金额','fn':'sum'}]}])['results'][0]
    assert result['sheet'] == '*'
    assert result['sources'] == [{'sheet':'支行01','range':[1,20],'rows_scanned':20},{'sheet':'支行02','range':[1,20],'rows_scanned':20}]
    assert result['filter'] is None
