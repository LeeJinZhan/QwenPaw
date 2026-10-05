from pathlib import Path
import json

import pytest
from docx import Document
from pptx import Presentation


@pytest.mark.parametrize('max_images,exported,limited', [(200,200,True), (250,201,False)])
def test_image_export_count_limit_preserves_native_text_and_reports_gap(tmp_path, monkeypatch, max_images, exported, limited):
    import io
    from types import SimpleNamespace
    from PIL import Image
    import docvortex
    from bank_mineru_mcp.native_office import extract_native_office
    source = tmp_path/'many-images.docx'
    document = Document()
    document.add_paragraph('原生正文保留')
    document.save(source)
    encoded = io.BytesIO()
    Image.new('RGB',(1,1),'white').save(encoded,format='PNG')
    parsed = SimpleNamespace(middle_json=SimpleNamespace(pages=[None]),diagnostics=[],
        assets={f'image-{index}':encoded.getvalue() for index in range(201)})
    monkeypatch.setattr(docvortex,'parse',lambda *args, **kwargs:parsed)
    monkeypatch.setattr(docvortex,'render_artifact',lambda *args, **kwargs:
        SimpleNamespace(content='原生正文保留'.encode(),diagnostics=[]))
    result = extract_native_office(source,tmp_path/'work',max_bytes=1024**2,max_images=max_images)
    assert '原生正文保留' in result['markdown']
    assert len(result['image_assets']) == exported
    assert ('native_image_count_limit' in result['diagnostics']) == limited
    assert result['image_export']['max_images'] == max_images
    assert result['image_export']['candidates_total'] == 201
    assert result['coverage']['image_text'] == 'unread'


def test_image_byte_quota_preserves_native_text(tmp_path, monkeypatch):
    import io
    import os
    from types import SimpleNamespace
    from PIL import Image
    import docvortex
    from bank_mineru_mcp.native_office import extract_native_office
    source = tmp_path/'image.doc'
    source.write_bytes(b'legacy-source')
    encoded = io.BytesIO()
    Image.frombytes('RGB',(32,32),os.urandom(32*32*3)).save(encoded,format='PNG')
    parsed = SimpleNamespace(middle_json=SimpleNamespace(pages=[None]),diagnostics=[],assets={'image':encoded.getvalue()})
    monkeypatch.setattr(docvortex,'parse',lambda *args, **kwargs:parsed)
    monkeypatch.setattr(docvortex,'render_artifact',lambda *args, **kwargs:
        SimpleNamespace(content=b'kept text',diagnostics=[]))
    result = extract_native_office(source,tmp_path/'work',max_bytes=1024)
    assert result['markdown'] == 'kept text'
    assert result['image_assets'] == []
    assert 'native_image_bytes_limit' in result['diagnostics']
    assert result['image_export']['skipped_byte_limit'] == 1


@pytest.mark.parametrize('warning', [False, True])
def test_native_pillow_bomb_skips_image_and_keeps_text(tmp_path, monkeypatch, warning):
    import warnings
    from types import SimpleNamespace
    from PIL import Image
    import docvortex
    from bank_mineru_mcp.native_office import extract_native_office
    source = tmp_path / 'bomb.doc'
    source.write_bytes(b'legacy-source')
    parsed = SimpleNamespace(middle_json=SimpleNamespace(pages=[None]), diagnostics=[],
                             assets={'image': b'oversized-image'})
    monkeypatch.setattr(docvortex, 'parse', lambda *args, **kwargs: parsed)
    monkeypatch.setattr(docvortex, 'render_artifact', lambda *args, **kwargs:
                        SimpleNamespace(content=b'kept native body', diagnostics=[]))
    def reject_image(*args, **kwargs):
        if warning:
            warnings.warn('too many pixels', Image.DecompressionBombWarning)
        raise Image.DecompressionBombError('too many pixels')
    monkeypatch.setattr(Image, 'open', reject_image)
    with warnings.catch_warnings(record=True) as emitted:
        warnings.simplefilter('always')
        result = extract_native_office(source, tmp_path / 'work', max_bytes=4096)
    assert not emitted
    assert result['markdown'] == 'kept native body'
    assert result['image_assets'] == []
    assert result['image_export']['skipped_pixel_limit'] == 1
    assert 'native_image_pixel_limit' in result['diagnostics']


def test_native_serializer_guard_is_restored_after_parser_failure(tmp_path, monkeypatch):
    from importlib import import_module
    import docvortex
    from bank_mineru_mcp.native_office import extract_native_office
    modules = [import_module('docvortex.analyzers.native.office.' + name)
               for name in ('doc.doc_converter', 'docx.resources', 'docx.tables',
                            'ppt.parser', 'pptx.resources', 'rtf.converter')]
    originals = [module.serialize_office_image for module in modules]
    def fail_parse(*args, **kwargs):
        raise RuntimeError('native engine failed')
    monkeypatch.setattr(docvortex, 'parse', fail_parse)
    source = tmp_path / 'failed.doc'
    source.write_bytes(b'legacy-source')
    with pytest.raises(RuntimeError, match='native engine failed'):
        extract_native_office(source, tmp_path / 'work', max_bytes=4096)
    assert [module.serialize_office_image for module in modules] == originals


def test_docx_native_parser_preserves_nested_tables_and_unicode(tmp_path):
    from bank_mineru_mcp.native_office import extract_native_office
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("银行原生内容 001234567890123456789")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "金额"
    table.cell(0, 1).text = "123.45"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套表格证据"
    document.save(source)
    result = extract_native_office(source, tmp_path / "work", max_bytes=1024 * 1024)
    assert "001234567890123456789" in result["markdown"]
    assert "嵌套表格证据" in result["markdown"]
    assert result["engine"] == "docvortex-0.5.4"
    assert result["source_inventory"]["tables"] == 2


@pytest.mark.parametrize('extension', ['.docx', '.pptx'])
def test_office_real_oversized_png_header_keeps_body(tmp_path, extension):
    import io
    import struct
    import zipfile
    import zlib
    from PIL import Image
    from bank_mineru_mcp.native_office import extract_native_office
    source = tmp_path / ('large-image' + extension)
    picture = io.BytesIO()
    Image.new('RGB', (8, 8), 'white').save(picture, format='PNG')
    picture.seek(0)
    if extension == '.docx':
        document = Document()
        document.add_paragraph('正文证据仍保留')
        document.add_picture(picture)
        document.save(source)
    else:
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = '正文证据仍保留'
        slide.shapes.add_picture(picture, Inches(1), Inches(2), width=Inches(2), height=Inches(2))
        presentation.save(source)
    with zipfile.ZipFile(source) as archive:
        entries = {info.filename: archive.read(info) for info in archive.infolist()}
    name = next(name for name in entries if name.startswith(('word/media/', 'ppt/media/')))
    payload = bytearray(entries[name])
    payload[16:24] = struct.pack('>II', 20000, 20000)
    payload[29:33] = struct.pack('>I', zlib.crc32(payload[12:29]) & 0xffffffff)
    entries[name] = bytes(payload)
    with zipfile.ZipFile(source, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    result = extract_native_office(source, tmp_path / 'work', max_bytes=1024**2)
    assert '正文证据仍保留' in result['markdown']
    assert not result['image_assets']
    assert result['image_export']['skipped_pixel_limit'] == 1, result


def test_pptx_native_parser_keeps_slide_text_and_notes(tmp_path):
    from bank_mineru_mcp.native_office import extract_native_office
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "复杂资料"
    slide.placeholders[1].text = "正文数据 100.25"
    slide.notes_slide.notes_text_frame.text = "独立演讲备注"
    presentation.save(source)
    result = extract_native_office(source, tmp_path / "work", max_bytes=1024 * 1024)
    assert "复杂资料" in result["markdown"]
    assert "正文数据" in result["markdown"]
    assert result["source_inventory"]["slides"] == 1
    assert result["source_inventory"]["notes"] == 1
    assert "独立演讲备注" in result["markdown"]
