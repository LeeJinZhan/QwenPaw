from datetime import datetime, timedelta, timezone
import json
from types import SimpleNamespace

import pytest


def test_xls_native_facts_keep_identifiers_hidden_sheet_and_raw_merge(tmp_path):
    import xlwt
    from bank_mineru_mcp.spreadsheet import extract_workbook
    book = xlwt.Workbook()
    sheet = book.add_sheet("旧表")
    sheet.write(0, 0, "编号")
    sheet.write(0, 1, "金额")
    sheet.write(1, 0, "001234567890123456789")
    sheet.write_merge(1, 2, 1, 1, 100.25)
    hidden = book.add_sheet("隐藏表")
    hidden.visibility = 1
    hidden.write(0, 0, "内部字段")
    source = tmp_path / "source.xls"
    book.save(source)
    work = tmp_path / "work"
    inventory = extract_workbook(source, work, stem="source")
    facts = [json.loads(line) for line in (work / inventory["sheets"][0]["source_file"]).read_text().splitlines()]
    assert facts[1]["v"] == ["001234567890123456789", 100.25]
    assert facts[2]["v"] == [None, None]
    assert inventory["sheets"][1]["hidden"] is True
    assert inventory["sheets"][0]["formula_expression_status"] == "unavailable"
from openpyxl import Workbook

from bank_mineru_mcp.spreadsheet import extract_workbook
from bank_mineru_mcp.structured_store import StructuredStore


def test_ragged_delimited_source_rows_keep_missing_cells_and_coordinates(tmp_path):
    task = tmp_path / 'task_ragged'
    task.mkdir()
    source = task / 'source.csv'
    source.write_text('id,amount,note\n001,100\n002\n', encoding='utf-8')
    work = task / 'work'
    inventory = extract_workbook(source, work, stem='source')
    store = StructuredStore(root=tmp_path)
    ref = store.write(SimpleNamespace(task_id='task_ragged', path=source,
        expires_at=datetime.now(timezone.utc) + timedelta(minutes=10)), inventory, work).document_ref
    result = store.read_range(ref, format='source')
    assert [row['values'] for row in result['records']] == [
        ['id','amount','note'], ['001','100',None], ['002',None,None]]
    assert [row['source_row'] for row in result['records']] == [1,2,3]


def store_workbook(tmp_path, rows, *, merge=None, header_row=1):
    task = tmp_path / "task_data"
    task.mkdir()
    source = task / "source.xlsx"
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    if merge:
        wb.active.merge_cells(merge)
    wb.save(source)
    inventory = extract_workbook(source, task / "work", stem="source", header_row=header_row)
    store = StructuredStore(root=tmp_path)
    handle = store.write(SimpleNamespace(task_id="task_data", path=source,
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1)), inventory, task / "work")
    return store, handle.document_ref, inventory


def test_raw_merge_values_match_read_and_compute(tmp_path):
    store, ref, inventory = store_workbook(tmp_path,
        [["id", "amount"], ["A", 100], ["B", None]], merge="B2:B3")
    result = store.read_range(ref, format="records")
    assert [row["values"][1] for row in result["records"]] == [100, None]
    assert sum(row["values"][1] or 0 for row in result["records"]) == store.aggregate(
        ref, [{"metrics": [{"column": "amount", "fn": "sum"}]}])["results"][0]["groups"][0]["amount:sum"]
    assert inventory["value_semantics"] == "raw_anchor"


def test_explicit_header_keeps_source_rows_and_coordinates(tmp_path):
    store, ref, inventory = store_workbook(tmp_path,
        [["Report", "Bank", "2026", "Unit"], ["id", "amount"], ["001", 10]], header_row=2)
    assert inventory["sheets"][0]["header_row"] == 2
    assert inventory["sheets"][0]["source_rows"] == 3
    assert store.read_range(ref, format="records")["records"] == [{"row": 1, "source_row": 3, "values": ["001", 10, None, None]}]
    facts = store.inventory(ref)["sheets"][0]["source_file"]
    entry, _ = store._entry(ref)
    assert len((entry.path / facts).read_text().splitlines()) == 3


def test_count_variants_and_explicit_numeric_text(tmp_path):
    store, ref, _ = store_workbook(tmp_path, [["id", "amount"], [1, "1,200"], [2, "800"], [3, None]])
    funcs = ["count", "count_rows", "count_nonempty", "count_numeric", "count_distinct", "sum"]
    result = store.aggregate(ref, [{"numeric_text": "thousands", "metrics": [{"column": "amount", "fn": fn} for fn in funcs]}])
    group = result["results"][0]["groups"][0]
    assert [group["amount:" + fn] for fn in funcs] == [3, 3, 2, 2, 2, 2000]
    assert result["results"][0]["semantics"]["numeric_text"] == "thousands"


def test_decimal_result_does_not_round_to_six_places(tmp_path):
    store, ref, _ = store_workbook(tmp_path, [["id", "amount"], [1, 0.0000001], [2, 0.0000002]])
    result = store.aggregate(ref, [{"metrics": [{"column": "amount", "fn": "sum"}]}])["results"][0]
    assert result["groups"][0]["amount:sum"] == pytest.approx(0.0000003)
    assert result["groups"][0]["exact_values"]["amount:sum"] == "0.0000003"


def test_source_metadata_keeps_formula_format_and_hidden_objects(tmp_path):
    source = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.append(["id", "formula"])
    book.active.append(["00012", "=1+2"])
    book.active["A2"].number_format = "00000"
    book.active.row_dimensions[2].hidden = True
    book.active.column_dimensions["B"].hidden = True
    book.save(source)
    inventory = extract_workbook(source, tmp_path / "facts", stem="source", allow_partial=True)
    sheet = inventory["sheets"][0]
    record = json.loads((tmp_path / "facts" / sheet["source_file"]).read_text().splitlines()[1])
    assert record["v"] == ["00012", None]
    assert record["cells"]["1"]["formula"] == "=1+2"
    assert record["cells"]["1"]["cache_status"] == "missing"
    assert record["cells"]["1"]["cached_type"] == "missing"
    assert record["cells"]["0"]["number_format"] == "00000"
    assert sheet["hidden_rows"] == [2] and sheet["hidden_columns"] == [2]


def test_source_range_reports_physical_columns_without_expanding_selection(tmp_path):
    store, ref, _ = store_workbook(tmp_path, [["id", "amount", "note"], ["001", 12.3, "private"]])
    result=store.read_range(ref,rows=[2,2],columns=['amount'],format='source')
    assert result['source_columns']==[2]
    assert result['records'][0]['values']==[12.3]
    assert not result['sheet_complete'] and not result['all_columns']


@pytest.mark.parametrize('cell_type,cached,cached_type,value,status',[
    ('n','2','n',2,'available'),('str','text','s','text','available'),
    ('b','1','b',True,'available'),('e','#DIV/0!','e','#DIV/0!','error'),
    ('n','','missing',None,'missing')])
def test_formula_source_cache_keeps_type_and_does_not_recalculate(tmp_path,cell_type,cached,cached_type,value,status):
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    source=tmp_path/'source.xlsx';book=Workbook();book.active.append(['formula']);book.active.append(['=1+1']);book.save(source)
    with ZipFile(source) as archive:parts={name:archive.read(name) for name in archive.namelist()}
    namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    document=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    cell=document.find(f'.//{namespace}c[@r="A2"]');cell.set('t',cell_type)
    cell.find(namespace+'v').text=cached
    parts['xl/worksheets/sheet1.xml']=ET.tostring(document)
    with ZipFile(source,'w') as archive:
        for name,data in parts.items():archive.writestr(name,data)
    inventory=extract_workbook(source,tmp_path/'facts',stem='source',allow_partial=True)
    row=json.loads((tmp_path/'facts'/inventory['sheets'][0]['source_file']).read_text().splitlines()[1])
    assert row['v']==[value]
    assert row['cells']['0']['formula']=='=1+1'
    assert row['cells']['0']['cached_type']==cached_type
    assert row['cells']['0']['cache_status']==status


def test_parse_cache_identity_changes_with_source_fact_contract(tmp_path):
    import hashlib,hmac
    from bank_mineru_mcp.parse_jobs import source_nonce
    source=SimpleNamespace(task_id='task_a',file_id='file_a',path=tmp_path/'source.xlsx',original_name='source.xlsx')
    source.path.write_bytes(b'fixture')
    store=SimpleNamespace(key=b'local-test-key')
    previous=json.dumps([source.task_id,source.file_id,hashlib.sha256(b'fixture').hexdigest(),
        source.original_name,'table-facts-3',1],sort_keys=True,separators=(',',':')).encode()
    assert source_nonce(store,source)!=hmac.new(store.key,previous,hashlib.sha256).digest()


def test_per_sheet_header_and_explicit_row_deduplication(tmp_path):
    task = tmp_path / "task_data"
    task.mkdir()
    source = task / "source.xlsx"
    book = Workbook()
    book.active.title = "A"
    for row in [["id", "amount"], ["x", 10], ["y", 20]]:
        book.active.append(row)
    sheet = book.create_sheet("B")
    for row in [["Report"], ["id", "amount"], ["x", 10], ["x", 30]]:
        sheet.append(row)
    book.save(source)
    inventory = extract_workbook(source, task / "work", stem="source", header_row={"A": 1, "B": 2})
    store = StructuredStore(root=tmp_path)
    ref = store.write(SimpleNamespace(task_id="task_data", path=source, expires_at=datetime.now(timezone.utc) + timedelta(hours=1)), inventory, task / "work").document_ref
    for mode, amount in [("append", 70), ("deduplicate_rows", 60)]:
        result = store.aggregate(ref, [{"cross_sheet_union": {"key_column": "id", "mode": mode}, "metrics": [{"column": "amount", "fn": "sum"}]}])["results"][0]
        assert result["groups"][0]["amount:sum"] == amount
        assert result["rows_scanned"] == 4
        assert result["union_mode"] == mode


def test_changed_fact_file_cannot_reuse_an_unchanged_manifest(tmp_path):
    store, ref, _ = store_workbook(tmp_path, [["id", "amount"], [1, 10]])
    store.inventory(ref)
    entry, manifest = store._entry(ref)
    path = entry.path / manifest["inventory"]["sheets"][0]["file"]
    path.write_text(path.read_text().replace("10", "99"))
    from bank_mineru_mcp.structured_store import StructuredStoreError
    with pytest.raises(StructuredStoreError, match="integrity"):
        store.inventory(ref)
