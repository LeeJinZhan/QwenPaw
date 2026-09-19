"""Structured spreadsheet extraction with merged-range expansion.

Structured formats (xlsx, csv, tsv) never go through MinerU layout parsing.
This module streams sheets into bounded row-block files plus an inventory so
read_range/aggregate/search can address sheets and row ranges directly.
"""

from __future__ import annotations

import codecs
import csv
import io
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

BLOCK_ROWS = 200
CHUNK_SIM_ROWS = 25


class SpreadsheetExtractError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def extract_workbook(path: Path, target_dir: Path, *, stem: str) -> dict[str, Any]:
    suffix = path.suffix.lower()
    target_dir.mkdir(parents=True, exist_ok=True)
    if suffix in {".csv", ".tsv"}:
        sheets = _extract_delimited(path, suffix)
    elif suffix == ".xlsx":
        sheets = _extract_xlsx(path)
    else:
        raise SpreadsheetExtractError(
            "FILE_TYPE_UNSUPPORTED", "Structured extraction supports xlsx/csv/tsv"
        )
    inventory_sheets = []
    for index, sheet in enumerate(sheets):
        filename = f"sheet_{index:02d}.rows.jsonl"
        offsets = _write_rows(target_dir / filename, sheet["rows"])
        inventory_sheets.append(
            {
                "name": sheet["name"],
                "index": index,
                "rows": len(sheet["rows"]),
                "cols": sheet["cols"],
                "header_row": sheet["header_row"],
                "formula_count": sheet.get("formula_count", 0),
                "formula_cache_status": "available" if sheet.get("formula_count") else "not_applicable",
                "merged_ranges": sheet["merged_ranges"],
                "columns": sheet["columns"],
                "file": filename,
                "block_rows": BLOCK_ROWS,
                "block_offsets": offsets,
                "legacy_blocks": _legacy_blocks(sheet),
            }
        )
    return {
        "engine": "ooxml-1" if suffix == ".xlsx" else "delimited-1",
        "title": path.name,
        "sheet_count": len(inventory_sheets),
        "total_rows": sum(sheet["rows"] for sheet in inventory_sheets),
        "sheets": inventory_sheets,
    }


def _write_rows(target: Path, rows: Iterable[tuple[int, list[Any]]]) -> list[int]:
    offsets: list[int] = []
    written = 0
    with target.open("wb") as handle:
        for position, (row_number, values) in enumerate(rows):
            if position % BLOCK_ROWS == 0:
                offsets.append(written)
            line = (
                json.dumps({"r": row_number, "v": values}, ensure_ascii=False) + "\n"
            ).encode("utf-8")
            handle.write(line)
            written += len(line)
    if not offsets:
        offsets.append(0)
    return offsets


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    text = str(value)
    return text


def _dtype(values: list[Any]) -> tuple[str, int]:
    nulls = sum(1 for value in values if value is None or value == "")
    seen = [value for value in values if value is not None and value != ""]
    kind = "str"
    if seen:
        if all(isinstance(value, bool) for value in seen):
            kind = "bool"
        elif all(isinstance(value, int) and not isinstance(value, bool) for value in seen):
            kind = "int"
        elif all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in seen):
            kind = "float"
        elif all(_is_dateish(value) for value in seen):
            kind = "date"
    return kind, nulls


def _is_dateish(value: Any) -> bool:
    return isinstance(value, str) and len(value) >= 8 and value[:4].isdigit()


def _profile(columns: list[str], rows: list[tuple[int, list[Any]]]) -> list[dict[str, Any]]:
    profile = []
    for index, name in enumerate(columns):
        values = [values[index] if index < len(values) else None for _, values in rows]
        kind, nulls = _dtype(values)
        profile.append({"name": name, "dtype": kind, "nulls": nulls})
    return profile


def _extract_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise SpreadsheetExtractError(
            "MINERU_UNAVAILABLE", "Spreadsheet extraction dependency is unavailable"
        ) from exc
    merged_by_sheet = _merged_ranges_from_xml(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    try:
        for position, worksheet in enumerate(workbook.worksheets):
            merged = merged_by_sheet.get(str(worksheet.title), [])
            raw: list[list[Any]] = []
            formula_count = 0
            for row, formula_row in zip(worksheet.iter_rows(), formulas.worksheets[position].iter_rows(), strict=True):
                for value, formula in zip(row, formula_row, strict=True):
                    if formula.data_type == "f":
                        formula_count += 1
                        if value.value is None or value.data_type == "e":
                            raise SpreadsheetExtractError("DOCUMENT_FORMULA_CACHE_MISSING", "Recalculate and save the workbook in Excel before uploading again")
                raw.append([_cell(value.value) for value in row])
            result = _finalize_sheet(str(worksheet.title), raw, merged)
            result["formula_count"] = formula_count
            sheets.append(result)
    finally:
        workbook.close()
        formulas.close()
    return sheets


def _merged_ranges_from_xml(path: Path) -> dict[str, list[list[int]]]:
    """Read mergeCell metadata straight from OOXML (read-only mode lacks it)."""
    import re
    import zipfile
    from xml.etree import ElementTree

    namespaces = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    result: dict[str, list[list[int]]] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            rels = ElementTree.fromstring(
                archive.read("xl/_rels/workbook.xml.rels")
            )
            rel_targets = {
                rel.get("Id"): rel.get("Target") for rel in rels
            }
            for sheet in workbook_xml.findall("main:sheets/main:sheet", namespaces):
                name = str(sheet.get("name") or "")
                target = rel_targets.get(sheet.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                ), "")
                if not target:
                    continue
                member = target if target.startswith("xl/") else "xl/" + target.lstrip("/")
                if target.startswith("/"):
                    member = target.lstrip("/")
                try:
                    sheet_xml = ElementTree.fromstring(archive.read(member))
                except KeyError:
                    continue
                ranges = []
                for merge in sheet_xml.findall(".//main:mergeCells/main:mergeCell", namespaces):
                    ref = str(merge.get("ref") or "")
                    match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", ref)
                    if not match:
                        continue
                    ranges.append(
                        [
                            int(match.group(2)),
                            _col_index(match.group(1)),
                            int(match.group(4)),
                            _col_index(match.group(3)),
                        ]
                    )
                result[name] = ranges
    except (OSError, zipfile.BadZipFile, ElementTree.ParseError):
        return {}
    return result


def _col_index(label: str) -> int:
    index = 0
    for char in label:
        index = index * 26 + (ord(char) - 64)
    return index


def _expand_merged(raw: list[list[Any]], merged: list[list[int]]) -> list[list[Any]]:
    if not merged:
        return raw
    width = max((len(row) for row in raw), default=0)
    expanded = [list(row) + [None] * (width - len(row)) for row in raw]
    anchors: dict[tuple[int, int, int, int], Any] = {}
    for min_row, min_col, max_row, max_col in merged:
        if 1 <= min_row <= len(expanded):
            row = expanded[min_row - 1]
            anchor = row[min_col - 1] if min_col - 1 < len(row) else None
            anchors[(min_row, min_col, max_row, max_col)] = anchor
    for row_index, row in enumerate(expanded, start=1):
        for min_row, min_col, max_row, max_col in merged:
            if min_row <= row_index <= max_row:
                value = anchors[(min_row, min_col, max_row, max_col)]
                for column in range(min_col, max_col + 1):
                    if column - 1 < len(row):
                        row[column - 1] = value
    return expanded


def _finalize_sheet(name: str, raw: list[list[Any]], merged: list[list[int]]) -> dict[str, Any]:
    rows = _expand_merged(raw, merged)
    header_row = 0
    for index, row in enumerate(rows, start=1):
        present = [value for value in row if value not in (None, "")]
        # A merged title band expands to identical values; headers must differ.
        if len(present) >= 2 and len(set(map(str, present))) >= 2:
            header_row = index
            break
    if header_row == 0:
        header_row = 1
    width = max((len(row) for row in rows), default=0)
    header = rows[header_row - 1] if rows else []
    originals = [str(header[index]).strip() if index < len(header) and header[index] is not None else ""
                 for index in range(width)]
    reserved = {name for name in originals if name}
    used: set[str] = set()
    columns: list[str] = []
    for index, original in enumerate(originals):
        base = original or f"col{index + 1}"
        column_name = base
        suffix = 2
        while column_name in used or (not original and column_name in reserved):
            column_name = f"{base}#{suffix}"
            suffix += 1
            while column_name in reserved:
                column_name = f"{base}#{suffix}"
                suffix += 1
        used.add(column_name)
        columns.append(column_name)
    data = [
        (number, row)
        for number, row in enumerate(rows[header_row:], start=1)
    ]
    return {
        "name": name,
        "cols": width,
        "header_row": header_row,
        "merged_ranges": merged,
        "columns": [{**column, "source_index": index, "original_name": originals[index]}
                    for index, column in enumerate(_profile(columns, data))],
        "rows": data,
    }


def _extract_delimited(path: Path, suffix: str) -> list[dict[str, Any]]:
    text = _read_text(path)
    delimiter = "\t" if suffix == ".tsv" else ","
    parsed = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    parsed = [[_cell(value) for value in row] for row in parsed]
    sheet = _finalize_sheet("data", parsed, [])
    return [sheet]


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    if data.startswith(codecs.BOM_UTF8):
        encodings = ("utf-8-sig",)
    elif data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        encodings = ("utf-16",)
    else:
        encodings = ("utf-8", "gb18030")
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SpreadsheetExtractError(
        "DOCUMENT_TEXT_ENCODING_UNSUPPORTED", "Spreadsheet text encoding is unsupported"
    )


def render_markdown(columns: list[str], rows: Iterable[tuple[int, list[Any]]]) -> str:
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    for _, values in rows:
        cells = ["" if value is None else str(value) for value in values]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


__all__ = [
    "BLOCK_ROWS",
    "CHUNK_SIM_ROWS",
    "SpreadsheetExtractError",
    "extract_workbook",
    "render_markdown",
]


def _legacy_blocks(sheet):
    """Stable row ranges, sized before MCP encoding; no text is discarded."""
    names = [column["name"] for column in sheet["columns"]]
    blocks = []
    start = 1
    rows = []
    for number, row in sheet["rows"]:
        candidate = rows + [(number, row)]
        size = len(json.dumps(render_markdown(names, candidate), ensure_ascii=False).encode("utf-8"))
        if rows and (len(rows) >= CHUNK_SIM_ROWS or size > 12000):
            blocks.append([start, number - 1])
            rows = []
            start = number
        rows.append((number, row))
    if rows:
        blocks.append([start, len(sheet["rows"])])
    return blocks
